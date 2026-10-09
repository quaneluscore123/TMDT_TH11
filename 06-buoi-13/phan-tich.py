"""Buoi 13 - NV2 (pheu theo nguon, chan doan) + NV3 (LTV, nguong chi, CPC) + diem cong (co mau A/B).

Chay tu thu muc goc repo:  python 06-buoi-13/phan-tich.py
Doc:  02-dataset/traffic_daily.csv, ad_spend.csv, customers_orders.csv, products.csv
Ghi:  06-buoi-13/phan-tich-buoi-13.xlsx  (moi bang trong bao cao deu co o day)  + in ra man hinh
"""
import csv
import math
import pathlib
import statistics as st
from collections import Counter, defaultdict

import openpyxl
from openpyxl.styles import Font, PatternFill

ROOT = pathlib.Path(__file__).resolve().parent.parent
DATA = ROOT / '02-dataset'
OUT = ROOT / '06-buoi-13' / 'phan-tich-buoi-13.xlsx'

STEPS = ['phien', 'xem_san_pham', 'them_gio', 'bat_dau_thanh_toan', 'don_hoan_tat']
STEP_NAMES = ['Phiên', 'Xem SP', 'Thêm giỏ', 'Thanh toán', 'Hoàn tất']
SOURCES = ['organic', 'paid_search', 'social_paid', 'social_organic', 'email']


def read(name):
    with open(DATA / name, encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))


T, A, C, P = read('traffic_daily.csv'), read('ad_spend.csv'), read('customers_orders.csv'), read('products.csv')
DAYS = len({r['ngay'] for r in T})
sheets = {}  # ten sheet -> (tieu de cot, cac dong)


def show(title, header, rows, sheet=None):
    print(f'\n## {title}')
    print(' | '.join(header))
    for r in rows:
        print(' | '.join(f'{v:,.4f}' if isinstance(v, float) and abs(v) < 1 else (f'{v:,.0f}' if isinstance(v, (int, float)) else str(v)) for v in r))
    if sheet:
        sheets.setdefault(sheet, []).append((title, header, rows))


# ---------------------------------------------------------------- NV2: pheu theo nguon
funnel = defaultdict(Counter)
for r in T:
    for k in STEPS + ['doanh_thu']:
        funnel[r['nguon']][k] += int(r[k])
        funnel['TỔNG'][k] += int(r[k])

rows = [[s] + [funnel[s][k] for k in STEPS] + [funnel[s]['doanh_thu']] for s in SOURCES + ['TỔNG']]
show(f'Bảng phễu {DAYS} ngày (01/06–30/07/2026) — số tuyệt đối', ['Nguồn'] + STEP_NAMES + ['Doanh thu'], rows, 'NV2-Pheu')

rate = {s: [funnel[s][STEPS[i + 1]] / funnel[s][STEPS[i]] for i in range(4)] for s in SOURCES + ['TỔNG']}
full = {s: funnel[s]['don_hoan_tat'] / funnel[s]['phien'] for s in SOURCES + ['TỔNG']}
rows = [[s] + rate[s] + [full[s]] for s in SOURCES + ['TỔNG']]
show('Tỷ lệ chuyển tiếp từng bước (chẩn đoán) và toàn phễu (chỉ để so sánh)',
     ['Nguồn', 'Phiên→Xem SP', 'Xem SP→Giỏ', 'Giỏ→Thanh toán', 'Thanh toán→Hoàn tất', 'Toàn phễu'], rows, 'NV2-Pheu')

# Do tin cay: khoang tin cay 95% cua tung ty le, va do dao dong theo ngay / theo thang
rows = []
for s in SOURCES:
    for i in range(4):
        n, p = funnel[s][STEPS[i]], rate[s][i]
        half = 1.96 * math.sqrt(p * (1 - p) / n)
        rows.append([s, f'{STEP_NAMES[i]}→{STEP_NAMES[i + 1]}', n, p, p - half, p + half])
show('Khoảng tin cậy 95% của từng tỷ lệ (cả 60 ngày)', ['Nguồn', 'Bước', 'Số vào bước', 'Tỷ lệ', 'Cận dưới', 'Cận trên'], rows, 'NV2-Do-tin-cay')

rows = []
for s in SOURCES:
    days = [r for r in T if r['nguon'] == s]
    sess = [int(r['phien']) for r in days]
    comp = [int(r['don_hoan_tat']) / int(r['bat_dau_thanh_toan']) for r in days if int(r['bat_dau_thanh_toan'])]
    months = []
    for m in ('2026-06', '2026-07'):
        c = Counter()
        for r in days:
            if r['ngay'].startswith(m):
                for k in STEPS:
                    c[k] += int(r[k])
        months.append(c['don_hoan_tat'] / c['bat_dau_thanh_toan'])
    rows.append([s, st.mean(sess), min(sess), max(sess), min(comp), max(comp), months[0], months[1]])
show('Dao động theo ngày và theo tháng của bước Thanh toán→Hoàn tất',
     ['Nguồn', 'Phiên/ngày TB', 'Phiên/ngày thấp nhất', 'cao nhất', 'TT→HT ngày thấp nhất', 'cao nhất', 'Tháng 6', 'Tháng 7'], rows, 'NV2-Do-tin-cay')

# Chon diem can thiep: so don tang them = so nguoi vao buoc x muc cai thien x cac ty le phia sau
def extra_orders(src, step, uplift):
    after = math.prod(rate[src][step + 1:])
    return funnel[src][STEPS[step]] * uplift * after


rows = []
for src in SOURCES + ['TỔNG']:
    for i in range(4):
        rows.append([src, f'{STEP_NAMES[i]}→{STEP_NAMES[i + 1]}', funnel[src][STEPS[i]], math.prod(rate[src][i + 1:]), extra_orders(src, i, 0.01)])
show('Độ nhạy: số đơn tăng thêm (60 ngày) cho mỗi +1 điểm % ở từng bước',
     ['Nguồn', 'Bước', 'Số người vào bước', 'Tích các tỷ lệ phía sau', 'Đơn tăng thêm / +1 điểm %'], rows, 'NV2-Can-thiep')

best_view = st.mean(rate[s][0] for s in ('organic', 'paid_search', 'social_organic'))
opts = [
    # ten, nguon, buoc, muc cai thien, can cu
    ('A. Trang đích quảng cáo MXH', 'social_paid', 0, (best_view - rate['social_paid'][0]) / 2,
     f"thu hẹp 1/2 khoảng cách tới TB 3 nguồn không phải email ({best_view:.1%})"),
    ('B. Hoàn tất thanh toán — social_paid', 'social_paid', 3, (rate['organic'][3] - rate['social_paid'][3]) / 2,
     f"thu hẹp 1/2 khoảng cách tới organic ({rate['organic'][3]:.1%})"),
    ('C. Hoàn tất thanh toán — social_organic', 'social_organic', 3, (rate['organic'][3] - rate['social_organic'][3]) / 2,
     f"thu hẹp 1/2 khoảng cách tới organic ({rate['organic'][3]:.1%})"),
    ('D. Xem SP→Thêm giỏ toàn site', 'TỔNG', 1, 0.03, 'giả định +3 điểm % (không có nguồn nào làm tốt hơn để so)'),
    ('E. Giỏ→Thanh toán toàn site', 'TỔNG', 2, 0.03, 'giả định +3 điểm % (không có nguồn nào làm tốt hơn để so)'),
]
aov = funnel['TỔNG']['doanh_thu'] / funnel['TỔNG']['don_hoan_tat']
rows, opt_result = [], {}
for name, src, i, up, basis in opts:
    n_in, after = funnel[src][STEPS[i]], math.prod(rate[src][i + 1:])
    extra = n_in * up * after
    month = extra * 30 / DAYS
    opt_result[name] = extra
    rows.append([name, src, f'{STEP_NAMES[i]}→{STEP_NAMES[i + 1]}', n_in, rate[src][i], up, after, extra, month,
                 month / (funnel['TỔNG']['don_hoan_tat'] * 30 / DAYS), basis])
show('So sánh phương án can thiệp: đơn tăng thêm ≈ số người vào bước × mức cải thiện × các tỷ lệ phía sau',
     ['Phương án', 'Nguồn', 'Bước', 'Số vào bước (60 ngày)', 'Tỷ lệ hiện tại', 'Mức cải thiện', 'Tích tỷ lệ phía sau',
      'Đơn tăng thêm (60 ngày)', 'Đơn tăng thêm / tháng', '% so với đơn hiện tại', 'Căn cứ mức cải thiện'], rows, 'NV2-Can-thiep')

# ---------------------------------------------------------------- Do luong: nen tang quang cao vs website
ads = defaultdict(Counter)
for r in A:
    for k in ('hien_thi', 'luot_nhap', 'chi_phi', 'chuyen_doi', 'doanh_thu_quy_ket'):
        ads[r['kenh']][k] += int(r[k])
rows = []
for ch in ('paid_search', 'social_paid'):
    a, f = ads[ch], funnel[ch]
    rows.append([ch, a['luot_nhap'], f['phien'], f['phien'] / a['luot_nhap'], a['chuyen_doi'], f['don_hoan_tat'],
                 a['chuyen_doi'] / f['don_hoan_tat'], a['doanh_thu_quy_ket'] / a['chi_phi'], f['doanh_thu'] / a['chi_phi']])
show('Đối chiếu số liệu nền tảng quảng cáo với số đo trên website (60 ngày)',
     ['Kênh', 'Lượt nhấp (nền tảng)', 'Phiên (website)', 'Phiên / lượt nhấp', 'Chuyển đổi (nền tảng)', 'Đơn (website)',
      'Nền tảng / website', 'ROAS nền tảng báo', 'ROAS theo đơn website'], rows, 'NV3-Quang-cao')

camp = defaultdict(Counter)
for r in A:
    for k in ('hien_thi', 'luot_nhap', 'chi_phi', 'chuyen_doi', 'doanh_thu_quy_ket'):
        camp[(r['kenh'], r['chien_dich'])][k] += int(r[k])
rows = [[ch, cp, v['luot_nhap'], v['chi_phi'], v['chi_phi'] / v['luot_nhap'], v['chuyen_doi'] / v['luot_nhap']]
        for (ch, cp), v in camp.items()]
show('CPC thực tế theo chiến dịch', ['Kênh', 'Chiến dịch', 'Lượt nhấp', 'Chi phí', 'CPC thực tế', 'Tỷ lệ chuyển đổi nền tảng báo'], rows, 'NV3-Quang-cao')

# ---------------------------------------------------------------- NV3: LTV
orders = sum(int(r['so_don_12thang']) for r in C)
revenue = sum(int(r['doanh_thu_12thang']) for r in C)
aov_c = revenue / orders
freq = orders / len(C)

# Ty le gop co trong so doanh thu: ty le gop tung nhom hang x doanh thu nhom hang do
CAT = {'Bút viết': 'Bút viết', 'Vở & sổ': 'Vở & sổ', 'Dụng cụ': 'Dụng cụ học tập', 'Ba lô & túi': 'Ba lô & túi', 'Mỹ thuật': 'Mỹ thuật'}
cat_price = defaultdict(lambda: [0, 0, 0])
for p in P:
    c = CAT[p['danh_muc'].split(' > ')[0]]
    cat_price[c][0] += int(p['gia_ban'])
    cat_price[c][1] += int(p['gia_von'])
    cat_price[c][2] += 1
cat_rev = Counter()
for r in C:
    cat_rev[r['nhom_hang_chinh']] += int(r['doanh_thu_12thang'])
rows = []
for c, (gb, gv, n) in cat_price.items():
    m = 1 - gv / gb
    rows.append([c, n, m, cat_rev[c], cat_rev[c] / revenue, m * cat_rev[c] / revenue])
margin_w = sum(r[-1] for r in rows)
margin_simple = st.mean(1 - int(p['gia_von']) / int(p['gia_ban']) for p in P)
rows.append(['Bình quân có trọng số theo doanh thu', '', margin_w, revenue, 1.0, margin_w])
rows.append(['(Sai) Bình quân đơn giản 40 sản phẩm', 40, margin_simple, '', '', ''])
show('Tỷ lệ lợi nhuận gộp có trọng số theo doanh thu',
     ['Nhóm hàng', 'Số SP', 'Tỷ lệ gộp nhóm', 'Doanh thu 12 tháng', 'Tỷ trọng doanh thu', 'Đóng góp'], rows, 'NV3-LTV')

repeat = sum(int(r['so_don_12thang']) > 1 for r in C) / len(C)
YEARS = 1.5  # gia dinh - xem can cu trong bao cao
ltv = aov_c * freq * YEARS * margin_w
cac_max = ltv / 3
rows = [['Đơn trung bình (AOV)', aov_c, f'{revenue:,} đ / {orders:,} đơn'],
        ['Số đơn / năm', freq, f'{orders:,} đơn / {len(C)} khách (12 tháng)'],
        ['Số năm gắn bó (GIẢ ĐỊNH)', YEARS, f'{repeat:.0%} khách mua > 1 đơn; chưa quan sát được quá 12 tháng'],
        ['Tỷ lệ lợi nhuận gộp (trọng số doanh thu)', margin_w, 'bảng trên'],
        ['LTV (lợi nhuận gộp)', ltv, 'AOV × đơn/năm × số năm × tỷ lệ gộp'],
        ['Ngưỡng chi thu hút tối đa (LTV/CAC = 3)', cac_max, 'LTV / 3']]
show('Giá trị vòng đời và ngưỡng chi thu hút', ['Chỉ số', 'Giá trị', 'Cách tính / nguồn'], rows, 'NV3-LTV')

rows = [[y, aov_c * freq * y * margin_w, aov_c * freq * y * margin_w / 3] for y in (1.0, 1.5, 2.0, 3.0)]
show('Độ nhạy theo giả định số năm gắn bó', ['Số năm', 'LTV', 'Ngưỡng chi (LTV/3)'], rows, 'NV3-LTV')

rows = []
for ch in SOURCES:
    rs = [r for r in C if r['kenh_thu_hut'] == ch]
    o = sum(int(r['so_don_12thang']) for r in rs)
    rv = sum(int(r['doanh_thu_12thang']) for r in rs)
    l = rv / o * o / len(rs) * YEARS * margin_w
    rows.append([ch, len(rs), rv / o, o / len(rs), l, l / 3])
show('LTV theo kênh thu hút (cùng giả định 1,5 năm và tỷ lệ gộp)', ['Kênh', 'Số khách', 'AOV', 'Đơn/năm', 'LTV', 'Ngưỡng chi'], rows, 'NV3-LTV')

# Muc chi moi luot nhap = nguong chi x ty le chuyen doi da do (don / phien tren website)
rows = []
for ch in ('paid_search', 'social_paid'):
    cr = full[ch]
    max_cpc = cac_max * cr
    spend = ads[ch]['chi_phi']
    cost_per_session = spend / funnel[ch]['phien']
    cpc = spend / ads[ch]['luot_nhap']
    cac_real = spend / funnel[ch]['don_hoan_tat']
    rows.append([ch, cac_max, cr, max_cpc, cpc, cost_per_session, cac_real, ltv / cac_real])
show('Mức chi tối đa mỗi lượt nhấp so với thực tế',
     ['Kênh', 'Ngưỡng chi', 'Tỷ lệ chuyển đổi đã đo (đơn/phiên)', 'CPC tối đa', 'CPC nền tảng báo', 'Chi phí / phiên đo được',
      'Chi phí / đơn đo được', 'LTV / chi phí mỗi đơn'], rows, 'NV3-Quang-cao')

# ---------------------------------------------------------------- Diem cong: co mau A/B
p1 = (funnel['social_paid']['don_hoan_tat'] + funnel['social_organic']['don_hoan_tat']) / \
     (funnel['social_paid']['bat_dau_thanh_toan'] + funnel['social_organic']['bat_dau_thanh_toan'])
rows = []
for mde in (0.05, 0.08, 0.10):
    p2 = p1 + mde
    pbar = (p1 + p2) / 2
    n = (1.96 * math.sqrt(2 * pbar * (1 - pbar)) + 0.8416 * math.sqrt(p1 * (1 - p1) + p2 * (1 - p2))) ** 2 / mde ** 2
    per_day = (funnel['social_paid']['bat_dau_thanh_toan'] + funnel['social_organic']['bat_dau_thanh_toan']) / DAYS
    rows.append([p1, p2, mde, math.ceil(n), math.ceil(2 * n / per_day), per_day])
show('Cỡ mẫu A/B (α = 0,05 hai phía, power = 80%) — đơn vị: lượt bắt đầu thanh toán từ MXH',
     ['Tỷ lệ hiện tại', 'Tỷ lệ mục tiêu', 'Mức cải thiện tối thiểu', 'Cỡ mẫu mỗi nhóm', 'Số ngày cần', 'Lượt TT từ MXH / ngày'], rows, 'Diem-cong-AB')

# ---------------------------------------------------------------- Xuat Excel
wb = openpyxl.Workbook()
wb.remove(wb.active)
bold, head_fill = Font(bold=True), PatternFill('solid', fgColor='D9EAD3')
for name, blocks in sheets.items():
    ws = wb.create_sheet(name)
    row = 1
    for title, header, data in blocks:
        ws.cell(row, 1, title).font = Font(bold=True, size=12)
        row += 1
        for j, h in enumerate(header, 1):
            c = ws.cell(row, j, h)
            c.font, c.fill = bold, head_fill
        for d in data:
            row += 1
            for j, v in enumerate(d, 1):
                c = ws.cell(row, j, v)
                if isinstance(v, float) and abs(v) < 1:
                    c.number_format = '0.0%'
                elif isinstance(v, (int, float)):
                    c.number_format = '#,##0'
        row += 2
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(12, min(45, max(len(str(c.value or '')) for c in col[1:]) + 2))
wb.save(OUT)
print(f'\n-> {OUT.relative_to(ROOT)}')
